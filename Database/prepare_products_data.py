"""Read ProductsData without changing it and build a validated import manifest.

Requires openpyxl and Pillow. No spreadsheet formulas or macros are executed.
"""
from collections import Counter
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import json
import re
import struct
import zipfile
import zlib

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    "cpu": "CPU/CPU.xlsx", "gpu": "Video_Card/GPU.xlsx",
    "mb": "Motherboard/Motherboard.xlsx", "memory": "Memory/RAM.xlsx",
    "storage": "Storage/ROM.xlsx", "psu": "Power_Supply/PSU.xlsx",
    "case_box": "Case/Case.xlsx", "cooling": "CPU_Cooler/CPU_Cooler.xlsx",
    "monitor": "Monitor/Monitor.xlsx",
}
RAM_IMAGES = {
    1: "G.SKILL Trident Z5 RGB", 2: "Corsair Dominator Titanium RGB",
    3: "TeamGroup T-Force Delta RGB", 4: "G.SKILL Trident Z5 Neo RGB (AMD EXPO)",
    5: "Corsair Vengeance RGB", 6: "G.SKILL Trident Z5 RGB",
    7: "Kingston Fury Renegade RGB", 8: "Corsair Dominator Platinum RGB",
    9: "TeamGroup T-Force XTREEM ARGB", 10: "G.SKILL Ripjaws S5",
    11: "Corsair Vengeance AMD EXPO", 12: "Patriot Viper Xtreme 5 RGB",
    13: "ADATA XPG Lancer RGB", 14: "Thermaltake TOUGHRAM XG RGB D5",
    15: "Crucial Pro Overclocking",
}
BRANDS = ["Fractal Design", "Cooler Master", "Super Flower", "be quiet!", "Lian Li",
          "G.SKILL", "TeamGroup", "Thermaltake", "Thermalright", "PowerColor",
          "SilverStone", "GIGABYTE", "Gigabyte", "Seasonic", "Solidigm", "Sapphire",
          "Kingston", "Samsung", "Seagate", "Phanteks", "Corsair", "Crucial",
          "Patriot", "Sabrent", "Noctua", "DeepCool", "ARCTIC", "ADATA", "NVIDIA",
          "Intel", "ASRock", "ASUS", "AMD", "NZXT", "EVGA", "MSI", "EKWB",
          "HYTE", "SSUPD", "FormD", "WD_BLACK", "WD", "Lexar", "Dell", "LG",
          "BenQ", "Acer", "Apple"]


def recover_zip(path):
    """Recover complete local entries only; refuse bad CRCs or truncated data."""
    data = path.read_bytes()
    output = BytesIO()
    offset = 0
    names = set()
    with zipfile.ZipFile(output, "w") as recovered:
        while data[offset:offset + 4] == b"PK\x03\x04":
            header = struct.unpack_from("<4s5H3I2H", data, offset)
            _, _, flags, method, _, _, crc, size, raw_size, name_len, extra_len = header
            if flags & 9 or method not in (0, 8):
                raise ValueError("Unsupported ZIP entry in " + str(path))
            name = data[offset + 30:offset + 30 + name_len].decode("utf-8")
            start = offset + 30 + name_len + extra_len
            compressed = data[start:start + size]
            raw = zlib.decompress(compressed, -15) if method == 8 else compressed
            if len(compressed) != size or len(raw) != raw_size or zlib.crc32(raw) != crc:
                raise ValueError("Corrupt workbook entry: " + name)
            if name in names:
                raise ValueError("Duplicate workbook entry: " + name)
            names.add(name)
            recovered.writestr(name, raw)
            offset = start + size
    if data[offset:offset + 4] != b"PK\x01\x02" or "xl/worksheets/sheet1.xml" not in names:
        raise ValueError("Cannot safely recover workbook " + str(path))
    output.seek(0)
    return output


def number(value):
    match = re.search(r"\d+", str(value))
    if not match:
        raise ValueError("Expected a numeric specification: " + str(value))
    return int(match[0])


def normalize(value):
    return re.sub(r'[^a-z0-9]', '', value.lower())


def prepare():
    records, sources, warnings = [], [], []
    for category, relative in SOURCES.items():
        path = ROOT / "ProductsData" / relative
        source = path
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
        except zipfile.BadZipFile:
            source = recover_zip(path)
            workbook = load_workbook(source, read_only=True, data_only=True)
            warnings.append(f"{relative}: damaged ZIP index; recovered all complete entries with CRC validation in memory. Original unchanged.")
        if len(workbook.worksheets) != 1:
            raise ValueError("Review new workbook layout: " + relative)
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        headers = rows[6]
        if headers[0] != "ID" or headers[-1] != "Price (USD)":
            raise ValueError("Unexpected headers: " + relative)
        products = [(row_num, row) for row_num, row in enumerate(rows, 1)
                    if isinstance(row[0], (int, float))]
        if len(products) != 15 or {r[0] for _, r in products} != set(range(1, 16)):
            raise ValueError("Expected 15 distinct product IDs: " + relative)
        sources.append({"path": path.relative_to(ROOT).as_posix(), "sha256": sha256(path.read_bytes()).hexdigest(), "sheet": sheet.title, "count": len(products)})
        images = [p for p in path.parent.rglob("*") if p.suffix.lower() in (".jpg", ".jpeg", ".png", ".avif", ".webp")]
        for row_num, row in products:
            source_id, name = int(row[0]), row[1].strip()
            if any(value is None for value in row):
                raise ValueError(f"Missing source cell: {relative}:{row_num}")
            price = Decimal(str(row[-1]))
            if not price.is_finite() or price <= 0 or price != price.quantize(Decimal("0.01")):
                raise ValueError("Invalid source price: " + name)
            matches = [p for p in images if normalize(p.stem) == normalize(RAM_IMAGES[source_id] if category == "memory" else name)]
            if len(matches) != 1:
                raise ValueError(f"Image match is not unique: {name}: {matches}")
            image = matches[0]
            with Image.open(image) as img:
                img.verify()
            brand = next((b for b in BRANDS if name.startswith(b + " ")), None)
            if brand is None:
                raise ValueError("Unknown brand: " + name)
            if brand == "Gigabyte": brand = "GIGABYTE"
            if brand == "WD_BLACK": brand = "WD"
            values = {"name": name, "brand": brand, "price": str(price.quantize(Decimal("0.01"))),
                      "image_url": f"productsdata-{category}-{source_id}{image.suffix.lower()}",
                      "description": "\n".join(f"{label}: {value}" for label, value in zip(headers[2:-1], row[2:-1]))}
            support = []
            if category == "cpu":
                cores = re.match(r"(.*?)\s*/\s*(\d+) Threads", row[3])
                if not cores: raise ValueError("Unknown CPU core layout: " + name)
                values.update(short_name=name, socket=row[2].replace(" ", ""), cores=cores[1], threads=cores[2],
                              base_clock=row[4].split(" (")[0], tdp=number(row[6]))
            elif category == "gpu":
                values.update(short_name=name, vram=row[3], boost_clock=row[5], tdp=number(row[6]), power_consumption=row[6])
            elif category == "mb":
                chipset = re.search(r"\b(Z790|X670E)\b", name)
                values.update(socket=row[2].replace(" ", ""), size=row[3], ram_slots=int(row[5]), memory_type=row[6], chipset=chipset[0] if chipset else None)
            elif category == "memory":
                kit = re.fullmatch(r"(\d+ GB) \((\d+)x(\d+)GB\)", row[4])
                if not kit: raise ValueError("Unknown memory kit: " + name)
                values.update(type=row[2], speed=row[3], capacity=kit[1], latency=row[5], modules=f"{kit[2]} x {kit[3]} GB", module_count=int(kit[2]))
            elif category == "storage":
                values.update(type="HDD" if "HDD" in name else "SSD", interface="NVMe" if "NVMe" in row[2] else "SATA",
                              capacity=row[4], read_speed=row[5], write_speed=row[6])
            elif category == "psu":
                values.update(wattage=row[2], wattage_watts=number(row[2]), rating=row[3], modularity=row[4])
            elif category == "case_box":
                values.update(size=row[2], side_panel=row[4], max_gpu_length=number(row[6]))
                support = [s.strip() for s in row[3].split(",")]
            elif category == "cooling":
                air = row[5] == "None (Air Cooler)"
                values.update(type="Air cooler" if air else "Liquid AIO cooler", size=None if air else f"{number(row[5])} mm", radiator_size_mm=None if air else number(row[5]))
                # This support is explicitly stated by the workbook subtitle (row 2).
                if "Compatible with LGA 1700 & AM5 Sockets" not in rows[1][0]:
                    raise ValueError("Review cooler socket support in source subtitle")
                support = ["LGA1700", "AM5"]
            elif category == "monitor":
                values.update(screen_size=row[2], panel_type=row[3], resolution=row[4], refresh_rate=row[5], response_time_hdr=row[6])
            records.append({"category": category, "source_id": source_id, "source_file": path.relative_to(ROOT).as_posix(),
                            "source_row": row_num, "source_values": dict(zip(headers, row)), "values": values,
                            "image_source": image.relative_to(ROOT).as_posix(), "image_sha256": sha256(image.read_bytes()).hexdigest(), "support": support})
        workbook.close()
    result = {"version": 1, "currency": "USD", "sources": sources, "warnings": warnings, "products": records}
    target = ROOT / "Database" / "products_data.json"
    target.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"products": len(records), "categories": dict(Counter(r["category"] for r in records)), "warnings": warnings}, indent=2))


if __name__ == "__main__":
    prepare()
